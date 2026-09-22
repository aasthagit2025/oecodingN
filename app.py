# -*- coding: utf-8 -*-
"""
OE Coder - single-file Streamlit app for open-end survey coding.

Everything lives in this one file, so it runs anywhere you can drop app.py
(Streamlit Cloud, a shared drive, a colleague's laptop) with no package layout
to get wrong. The oecoder/ package is the same code split into modules, if you
would rather import it from a script or a notebook.

Run with:   streamlit run app.py
"""
from __future__ import annotations

import json
import re
from collections import Counter
from io import BytesIO
from pathlib import Path
from typing import Iterable, Sequence

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# ============================================================================
#  ENGINE - coding logic
# ============================================================================

# --------------------------------------------------------------- housekeeping
HOUSEKEEPING = [
    {"net": "Generic / Non-specific", "code": 995,
     "label": "Generic positive / non-specific mention only - no specific reason given", "patterns": []},
    {"net": "Generic / Non-specific", "code": 996,
     "label": "Any other mention", "patterns": []},
    {"net": "Generic / Non-specific", "code": 997,
     "label": "Irrelevant / illegible / does not answer the question", "patterns": []},
    {"net": "Generic / Non-specific", "code": 998,
     "label": "Don't know / Can't say / Nothing / Not applicable", "patterns": []},
    {"net": "Generic / Non-specific", "code": 999,
     "label": "No response (blank)", "patterns": []},
]
HK_CODES = {r["code"] for r in HOUSEKEEPING}

# Fallback classifiers. Override per project if a study needs different wording.
DEFAULT_GENERIC_POS = (
    r"^(it'?s |its |this is |i think it is |overall |everthing is |everything is )?"
    r"(very+\s+|really\s+|so\s+|too\s+|highly\s*,?\s*|extremely\s+|quite\s+)*"
    r"(good|nice|great|excellent|excelent|exceleent|best|super|superb|awesome|awsome|awsem|awesom|fine|"
    r"ok+|okay|okk+|wonderful|amazing|fantastic|marvellous|outstanding|cool|perfect|happy|yes+|y|"
    r"agree|agri|average|strong|strongly|high|top|better|rigorously|bravely|lovely|verry good|noce|tes|prode)"
    r"([\s\-,]*(brand|company|companey|compnay|organi[sz]ation|organaisation|experience|experiance|"
    r"feel|feeling|see|one|and better|and great|guide|very|branch name|for me|absolutely|best|good))*"
    r"[\s\.\!\,\-\u2019'\*\u00a0\U0001F300-\U0001FAFF]*$"
)

DEFAULT_NULL = (
    r"^(na|n\.?a\.?|n/a|nil|no|none|nothing|not|no comments?|no comment|don'?t know|dont know|"
    r"can'?t say|cant say|not applicable|no idea|no response|no suggestions?|unable to say)"
    r"[\s\.\!\,\-]*$"
)

DEFAULT_JUNK = (
    r"^[\s\.\,\-\_\*\?/\\]*$|^[a-z]{1,2}[\s\.]*$|^\d+[\s\.\,\-/]*\d*[\s\.\,\-/]*\d*$|"
    r"^[b-df-hj-np-tv-z]{6,}$"
)


class Rules:
    """Compiled fallback rules for one question."""

    def __init__(self, generic_pos=DEFAULT_GENERIC_POS, null=DEFAULT_NULL, junk=DEFAULT_JUNK):
        self.generic_pos = re.compile(generic_pos, re.I)
        self.null = re.compile(null, re.I)
        self.junk = re.compile(junk, re.I)


# --------------------------------------------------------------- helpers
def clean(text) -> str:
    """Normalise a raw cell into a comparable verbatim string."""
    if text is None:
        return ""
    s = str(text)
    if s.strip().lower() in ("nan", "none", "null"):
        return ""
    return re.sub(r"\s+", " ", s).strip()


def compile_frame(frame: Sequence[dict]) -> list[tuple[int, list[re.Pattern]]]:
    """Pre-compile a frame once so a 10k-row file does not recompile per row."""
    out = []
    for row in frame:
        if int(row["code"]) in HK_CODES:
            continue  # housekeeping codes are assigned by fallback, never by pattern
        pats = [p for p in (row.get("patterns") or []) if str(p).strip()]
        out.append((int(row["code"]), [re.compile(p, re.I) for p in pats]))
    return out


def code_text(text, compiled, rules: Rules | None = None) -> list[int]:
    """Return the sorted list of codes for a single verbatim."""
    rules = rules or Rules()
    t = clean(text)
    if t == "":
        return [999]
    low = t.lower()

    hits = [code for code, pats in compiled if any(p.search(low) for p in pats)]
    if hits:
        return sorted(set(hits))

    if rules.null.match(t):
        return [998]
    if rules.generic_pos.match(t):
        return [995]
    if rules.junk.match(t) or len(re.sub(r"[^a-z]", "", low)) <= 2:
        return [997]
    return [996]


def code_series(values: Iterable, frame: Sequence[dict], rules: Rules | None = None) -> list[list[int]]:
    compiled = compile_frame(frame)
    rules = rules or Rules()
    return [code_text(v, compiled, rules) for v in values]


def apply_overrides(ids: Iterable, codes: list[list[int]], overrides: dict) -> list[list[int]]:
    """overrides: {respondent_id: [codes]}. Keys are matched as-is and as strings."""
    out = []
    for rid, cl in zip(ids, codes):
        if rid in overrides:
            out.append(sorted(set(int(c) for c in overrides[rid])))
        elif str(rid) in overrides:
            out.append(sorted(set(int(c) for c in overrides[str(rid)])))
        else:
            out.append(cl)
    return out


def code_dataframe(df: pd.DataFrame, id_col: str, verbatim_col: str,
                   frame: Sequence[dict], overrides: dict | None = None,
                   rules: Rules | None = None) -> pd.DataFrame:
    """Return a copy of df with normalised Verbatim and a Codes column."""
    out = df.copy()
    out["Verbatim"] = out[verbatim_col].map(clean)
    codes = code_series(out["Verbatim"], frame, rules)
    if overrides:
        codes = apply_overrides(out[id_col], codes, overrides)
    out["Codes"] = codes
    return out


# --------------------------------------------------------------- reporting
def frequency_table(coded: pd.DataFrame, frame: Sequence[dict], base_n: int | None = None) -> pd.DataFrame:
    """Counts and percentages for every code in the frame, in frame order."""
    from collections import Counter

    counts = Counter(c for cl in coded["Codes"] for c in cl)
    n = base_n or len(coded)
    answered = sum(1 for cl in coded["Codes"] if 999 not in cl)

    rows = []
    for row in list(frame) + [r for r in HOUSEKEEPING
                              if int(r["code"]) not in {int(x["code"]) for x in frame}]:
        code = int(row["code"])
        cnt = counts.get(code, 0)
        rows.append({
            "NET": row["net"],
            "Code": code,
            "Code label": row["label"],
            "Count": cnt,
            "% of Total": cnt / n if n else 0,
            "% of Answered": cnt / answered if answered else 0,
        })
    return pd.DataFrame(rows)


def unresolved(coded: pd.DataFrame, id_col: str) -> pd.DataFrame:
    """Rows that landed in 996/997 - the analyst review queue."""
    mask = coded["Codes"].map(lambda cl: bool({996, 997} & set(cl)))
    out = coded.loc[mask, [id_col, "Verbatim", "Codes"]].copy()
    out["Codes"] = out["Codes"].map(lambda cl: "/".join(map(str, cl)))
    return out


# ============================================================================
#  FRAMES - code frame load / save / validate
# ============================================================================

PATTERN_SEP = "\n"   # newline, NOT "|" - that character is regex alternation


# --------------------------------------------------------------- validation
def validate(frame: list[dict]) -> list[str]:
    """Return a list of human-readable problems. Empty list means the frame is clean."""
    problems = []
    seen = {}
    for i, row in enumerate(frame, start=1):
        for key in ("net", "code", "label"):
            if not str(row.get(key, "")).strip():
                problems.append(f"Row {i}: '{key}' is empty.")
        try:
            code = int(row["code"])
        except (KeyError, TypeError, ValueError):
            problems.append(f"Row {i}: code '{row.get('code')}' is not a whole number.")
            continue
        if code in seen:
            problems.append(f"Code {code} appears more than once (rows {seen[code]} and {i}).")
        seen[code] = i
        if code in HK_CODES:
            problems.append(f"Row {i}: {code} is a reserved housekeeping code - it is added automatically.")
        for p in row.get("patterns") or []:
            try:
                re.compile(p)
            except re.error as e:
                problems.append(f"Code {code}: pattern {p!r} is not valid regex ({e}).")
    return problems


def normalise(frame: list[dict]) -> list[dict]:
    """Coerce types and sort by code so the frame reads cleanly."""
    out = []
    for row in frame:
        pats = row.get("patterns") or []
        if isinstance(pats, str):
            pats = [p.strip() for p in pats.splitlines() if p.strip()]
        out.append({
            "net": str(row["net"]).strip(),
            "code": int(row["code"]),
            "label": str(row["label"]).strip(),
            "patterns": [str(p).strip() for p in pats if str(p).strip()],
        })
    return sorted(out, key=lambda r: r["code"])


# --------------------------------------------------------------- JSON
def save_json(frames: dict[str, list[dict]], path: str | Path, meta: dict | None = None) -> None:
    """frames: {question_key: frame}. Saved with an optional project meta block."""
    payload = {"meta": meta or {}, "frames": {q: normalise(f) for q, f in frames.items()}}
    Path(path).write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")


def load_json(path: str | Path) -> tuple[dict[str, list[dict]], dict]:
    payload = json.loads(Path(path).read_text(encoding="utf-8"))
    if "frames" in payload:                       # full project file
        return {q: normalise(f) for q, f in payload["frames"].items()}, payload.get("meta", {})
    if isinstance(payload, list):                 # bare single frame
        return {"Q1": normalise(payload)}, {}
    return {q: normalise(f) for q, f in payload.items()}, {}


# --------------------------------------------------------------- Excel round-trip
def to_excel(frames: dict[str, list[dict]], path: str | Path) -> None:
    """One sheet per question, editable by hand, re-readable by from_excel."""
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for q, frame in frames.items():
            rows = [{
                "NET": r["net"],
                "Code": r["code"],
                "Code label": r["label"],
                "Patterns (one per line)": PATTERN_SEP.join(r["patterns"]),
            } for r in normalise(frame)]
            pd.DataFrame(rows).to_excel(xw, sheet_name=q[:31], index=False)


def from_excel(path: str | Path) -> dict[str, list[dict]]:
    frames = {}
    book = pd.read_excel(path, sheet_name=None)
    for sheet, df in book.items():
        cols = {c.lower().strip(): c for c in df.columns}
        pat_col = next((cols[c] for c in cols if c.startswith("pattern")), None)
        rows = []
        for _, r in df.iterrows():
            if pd.isna(r.get(cols.get("code"))):
                continue
            pats = r.get(pat_col, "") if pat_col else ""
            pats = [] if pd.isna(pats) else [p.strip() for p in str(pats).splitlines() if p.strip()]
            rows.append({
                "net": r[cols["net"]],
                "code": r[cols["code"]],
                "label": r[cols.get("code label", cols.get("label"))],
                "patterns": pats,
            })
        if rows:
            frames[sheet] = normalise(rows)
    return frames


# --------------------------------------------------------------- misc
def with_housekeeping(frame: list[dict]) -> list[dict]:
    """Frame plus the five reserved codes, for display and reporting."""
    have = {int(r["code"]) for r in frame}
    return list(frame) + [r for r in HOUSEKEEPING if r["code"] not in have]


def blank_frame() -> list[dict]:
    return [{"net": "NET 1", "code": 101, "label": "First code", "patterns": ["example|sample"]}]


# ============================================================================
#  EXCEL - deliverable builder
# ============================================================================

ARIAL = "Arial"
H_FILL = PatternFill("solid", fgColor="1F3864")
NET_FILL = PatternFill("solid", fgColor="D9E2F3")
Q_FILL = PatternFill("solid", fgColor="8EA9DB")
HK_FILL = PatternFill("solid", fgColor="F2F2F2")
H_FONT = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
B_FONT = Font(name=ARIAL, size=10, bold=True)
N_FONT = Font(name=ARIAL, size=10)
_THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = H_FONT
        cell.fill = H_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def _widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def build_workbook(questions: dict, meta: dict, base_n: int | None = None) -> BytesIO:
    """
    questions: {q_key: {"title": str, "frame": [...], "coded": DataFrame,
                        "id_col": str, "extra_cols": [str, ...]}}
               coded must carry 'Verbatim' and 'Codes' columns.
    meta:      free-form project details shown on the Read Me sheet.
    Returns an in-memory .xlsx ready to write to disk or hand to a download button.
    """
    wb = Workbook()

    # ----------------------------------------------------------- Read Me
    ws = wb.active
    ws.title = "Read Me"
    ws["A1"] = meta.get("title", "Open-End Coding Deliverable")
    ws["A1"].font = Font(name=ARIAL, size=14, bold=True, color="1F3864")

    lines = list(meta.get("fields", {}).items())
    lines += [
        ("", ""),
        ("Coding approach", "Frame built from a read of the verbatims and applied as keyword / regex rules, "
                            "with analyst overrides on reviewed cases. Multi-coding is allowed - one verbatim "
                            "can carry more than one code."),
        ("Base for %", "Percentages are on Total N. An 'Answered' base excluding blanks (code 999) is also shown."),
        ("Housekeeping codes", "995 Generic non-specific | 996 Any other mention | 997 Irrelevant / illegible | "
                               "998 Don't know / Can't say / NA | 999 No response"),
        ("", ""),
        ("Sheet guide", ""),
        ("  Code Structure", "The code frame for every question with counts and percentages - the sheet to "
                             "circulate for client sign-off."),
        ("  <Q> Coded", "Respondent-level data: ID, verbatim and the codes assigned."),
        ("  <Q> Binary", "The same data as a 0/1 matrix, one column per code, for the tabulation setup."),
    ]
    r = 3
    for k, v in lines:
        ws.cell(row=r, column=1, value=k).font = B_FONT
        c = ws.cell(row=r, column=2, value=v)
        c.font = N_FONT
        c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    _widths(ws, {"A": 26, "B": 105})

    # ----------------------------------------------------------- Code Structure
    ws = wb.create_sheet("Code Structure")
    ws["A1"] = "CODE STRUCTURE - " + meta.get("title", "")
    ws["A1"].font = Font(name=ARIAL, size=13, bold=True, color="1F3864")
    ws["A2"] = "Multi-code frame | Codes 995-999 are housekeeping codes"
    ws["A2"].font = Font(name=ARIAL, size=9, italic=True)

    hdr = ["NET", "Code", "Code label", "Count", "% of Total", "% of Answered"]
    r = 4
    for q, spec in questions.items():
        coded = spec["coded"]
        n = base_n or len(coded)
        answered = sum(1 for cl in coded["Codes"] if 999 not in cl)
        counts = Counter(c for cl in coded["Codes"] for c in cl)

        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        cell = ws.cell(row=r, column=1, value=spec.get("title", q))
        cell.font = Font(name=ARIAL, size=11, bold=True, color="1F3864")
        cell.fill = Q_FILL
        ws.row_dimensions[r].height = 20
        r += 1
        ws.cell(row=r, column=1,
                value=f"Base: Total N = {n}   |   Answered = {answered}   |   No response = {n - answered}"
                ).font = Font(name=ARIAL, size=9, italic=True)
        r += 1

        for i, h in enumerate(hdr, start=1):
            ws.cell(row=r, column=i, value=h)
        _header(ws, r, 6)
        r += 1

        last_net = None
        for row in with_housekeeping(spec["frame"]):
            code = int(row["code"])
            if row["net"] != last_net:
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
                c = ws.cell(row=r, column=1, value=str(row["net"]).upper())
                c.font = Font(name=ARIAL, size=10, bold=True, color="1F3864")
                c.fill = NET_FILL
                c.border = BORDER
                r += 1
                last_net = row["net"]
            ws.cell(row=r, column=2, value=code)
            ws.cell(row=r, column=3, value=row["label"])
            ws.cell(row=r, column=4, value=counts.get(code, 0))
            ws.cell(row=r, column=5, value=f"=IF($D{r}=\"\",\"\",$D{r}/{n})")
            ws.cell(row=r, column=6, value=f"=IF($D{r}=\"\",\"\",$D{r}/{max(answered,1)})")
            for c_ in range(1, 7):
                cell = ws.cell(row=r, column=c_)
                cell.font = N_FONT
                cell.border = BORDER
                if code >= 995:
                    cell.fill = HK_FILL
            ws.cell(row=r, column=2).alignment = Alignment(horizontal="center")
            ws.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(row=r, column=4).alignment = Alignment(horizontal="center")
            ws.cell(row=r, column=5).number_format = "0.0%"
            ws.cell(row=r, column=6).number_format = "0.0%"
            r += 1
        r += 2

    _widths(ws, {"A": 34, "B": 8, "C": 78, "D": 9, "E": 13, "F": 15})
    ws.freeze_panes = "A5"

    # ----------------------------------------------------------- per question
    for q, spec in questions.items():
        coded = spec["coded"]
        frame = with_housekeeping(spec["frame"])
        labels = {int(x["code"]): x["label"] for x in frame}
        id_col = spec["id_col"]
        extra = [c for c in spec.get("extra_cols", []) if c in coded.columns]
        mx = max((len(cl) for cl in coded["Codes"]), default=1)

        ws = wb.create_sheet(f"{q} Coded"[:31])
        ws["A1"] = spec.get("title", q)
        ws["A1"].font = Font(name=ARIAL, size=11, bold=True, color="1F3864")
        cols = [id_col] + extra + ["Verbatim response"] + [f"Code_{i}" for i in range(1, mx + 1)] \
               + ["No. of codes", "Code labels"]
        for i, h in enumerate(cols, start=1):
            ws.cell(row=3, column=i, value=h)
        _header(ws, 3, len(cols))

        vcol = 2 + len(extra)
        r = 4
        for _, row in coded.iterrows():
            cl = row["Codes"]
            ws.cell(row=r, column=1, value=row[id_col])
            for j, ec in enumerate(extra):
                ws.cell(row=r, column=2 + j, value=row[ec])
            ws.cell(row=r, column=vcol, value=row["Verbatim"])
            for i in range(mx):
                ws.cell(row=r, column=vcol + 1 + i, value=cl[i] if i < len(cl) else None)
            ws.cell(row=r, column=vcol + 1 + mx, value=len(cl))
            ws.cell(row=r, column=vcol + 2 + mx,
                    value=" | ".join(labels.get(c, "") for c in cl))
            for c_ in range(1, len(cols) + 1):
                cell = ws.cell(row=r, column=c_)
                cell.font = N_FONT
                cell.border = BORDER
                is_code = vcol + 1 <= c_ <= vcol + 1 + mx
                cell.alignment = Alignment(vertical="top",
                                           wrap_text=c_ in (vcol, len(cols)),
                                           horizontal="center" if is_code else "left")
            r += 1

        spec_w = {"A": 10, get_column_letter(vcol): 62}
        for j in range(len(extra)):
            spec_w[get_column_letter(2 + j)] = 16
        for i in range(mx):
            spec_w[get_column_letter(vcol + 1 + i)] = 8
        spec_w[get_column_letter(vcol + 1 + mx)] = 11
        spec_w[get_column_letter(vcol + 2 + mx)] = 60
        _widths(ws, spec_w)
        ws.freeze_panes = ws.cell(row=4, column=vcol + 1).coordinate
        ws.auto_filter.ref = f"A3:{get_column_letter(len(cols))}{r - 1}"

        # ------------------------------------------------------- binary
        counts = Counter(c for cl in coded["Codes"] for c in cl)
        used = [int(x["code"]) for x in frame if counts.get(int(x["code"]), 0) > 0]
        ws = wb.create_sheet(f"{q} Binary"[:31])
        ws.cell(row=1, column=1, value=f"{q} binary matrix (1 = code applies)").font = B_FONT
        for i, c in enumerate(used):
            cell = ws.cell(row=1, column=3 + i, value=labels.get(c, ""))
            cell.font = Font(name=ARIAL, size=7)
            cell.alignment = Alignment(text_rotation=90, vertical="bottom")
        ws.row_dimensions[1].height = 165

        bh = [id_col, "Verbatim"] + [f"C{c}" for c in used]
        for i, h in enumerate(bh, start=1):
            ws.cell(row=2, column=i, value=h)
        _header(ws, 2, len(bh))

        r = 3
        for _, row in coded.iterrows():
            s = set(row["Codes"])
            ws.cell(row=r, column=1, value=row[id_col]).font = N_FONT
            ws.cell(row=r, column=2, value=row["Verbatim"][:120]).font = N_FONT
            for i, c in enumerate(used):
                cell = ws.cell(row=r, column=3 + i, value=1 if c in s else 0)
                cell.font = N_FONT
                cell.alignment = Alignment(horizontal="center")
            r += 1
        ws.cell(row=r, column=2, value="TOTAL").font = B_FONT
        for i in range(len(used)):
            col = get_column_letter(3 + i)
            cell = ws.cell(row=r, column=3 + i, value=f"=SUM({col}3:{col}{r - 1})")
            cell.font = B_FONT
            cell.alignment = Alignment(horizontal="center")

        bw = {"A": 10, "B": 45}
        for i in range(len(used)):
            bw[get_column_letter(3 + i)] = 6
        _widths(ws, bw)
        ws.freeze_panes = "C3"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ============================================================================
#  STREAMLIT APP
# ============================================================================

st.set_page_config(page_title="OE Coder", page_icon="🗂️", layout="wide")

PATTERN_HELP = ("One pattern per line. Each is a case-insensitive regular expression — plain "
                "words work fine (`brand name`), `\\b` marks a word boundary (`\\bpay\\b` will not "
                "match *payment*), and `(le|ility)` gives alternatives. Never separate patterns "
                "with `|`: that character is regex alternation and belongs inside a pattern.")


# ----------------------------------------------------------------- state
def init_state():
    st.session_state.setdefault("raw", {})          # q_key -> DataFrame
    st.session_state.setdefault("frames", {})       # q_key -> frame
    st.session_state.setdefault("titles", {})       # q_key -> question text
    st.session_state.setdefault("overrides", {})    # q_key -> {id: [codes]}
    st.session_state.setdefault("id_col", None)
    st.session_state.setdefault("extra_cols", [])
    st.session_state.setdefault("meta", {})


init_state()


def coded_for(q: str) -> pd.DataFrame:
    df = st.session_state.raw[q]
    return code_dataframe(
        df,
        id_col=st.session_state.id_col,
        verbatim_col=st.session_state.titles[q]["col"],
        frame=st.session_state.frames.get(q, []),
        overrides=st.session_state.overrides.get(q, {}),
    )


# ----------------------------------------------------------------- sidebar
with st.sidebar:
    st.title("🗂️ OE Coder")
    st.caption("Rule-based open-end coding with an analyst review step.")

    st.subheader("Project details")
    st.session_state.meta["title"] = st.text_input(
        "Deliverable title", st.session_state.meta.get("title", "Open-End Coding Deliverable"))
    job = st.text_input("Job number", st.session_state.meta.get("job", ""))
    client = st.text_input("Client", st.session_state.meta.get("client", ""))
    st.session_state.meta["job"] = job
    st.session_state.meta["client"] = client

    st.divider()
    st.subheader("Code frame library")
    lib = st.file_uploader("Load a saved frame (.json)", type=["json"], key="lib_json")
    if lib and st.button("Load frames", use_container_width=True):
        payload = json.loads(lib.getvalue().decode("utf-8"))
        loaded, meta = (payload.get("frames", payload), payload.get("meta", {})) \
            if isinstance(payload, dict) else ({"Q1": payload}, {})
        st.session_state.frames.update({q: normalise(f) for q, f in loaded.items()})
        st.success(f"Loaded {len(loaded)} frame(s): {', '.join(loaded)}")

    if st.session_state.frames:
        st.download_button(
            "💾 Save frames as .json",
            data=json.dumps({"meta": st.session_state.meta,
                             "frames": {q: normalise(f) for q, f in st.session_state.frames.items()}},
                            indent=2, ensure_ascii=False),
            file_name=f"{job or 'project'}_codeframes.json",
            mime="application/json", use_container_width=True)


# ----------------------------------------------------------------- tabs
tab_data, tab_frame, tab_run, tab_review, tab_export = st.tabs(
    ["1 · Data", "2 · Code frame", "3 · Run coding", "4 · Review", "5 · Export"])

# ============================================================ 1 · DATA
with tab_data:
    st.header("Load the open-end data")
    up = st.file_uploader("Excel or CSV with the verbatims", type=["xlsx", "xlsm", "csv"])

    if up:
        if up.name.lower().endswith(".csv"):
            book = {"data": pd.read_csv(up)}
        else:
            book = pd.read_excel(up, sheet_name=None)

        st.success(f"Found {len(book)} sheet(s): {', '.join(book)}")
        sheets = st.multiselect("Sheets to code (one question per sheet)",
                                list(book), default=list(book))

        if sheets:
            first = book[sheets[0]]
            st.session_state.id_col = st.selectbox(
                "Respondent ID column", list(first.columns),
                index=list(first.columns).index(st.session_state.id_col)
                if st.session_state.id_col in first.columns else 0)
            st.session_state.extra_cols = st.multiselect(
                "Other columns to carry through (password, cell, weight…)",
                [c for c in first.columns if c != st.session_state.id_col])

            st.markdown("**Pick the verbatim column on each sheet**")
            for sh in sheets:
                df = book[sh]
                guess = [c for c in df.columns
                         if c not in (st.session_state.id_col, *st.session_state.extra_cols)]
                col = st.selectbox(f"{sh}", guess, key=f"vcol_{sh}")
                st.session_state.raw[sh] = df
                st.session_state.titles[sh] = {"col": col, "text": str(col)}

            st.divider()
            st.markdown("**Question wording shown on the deliverable**")
            for sh in sheets:
                st.session_state.titles[sh]["text"] = st.text_input(
                    f"{sh} title", st.session_state.titles[sh]["text"], key=f"title_{sh}")

            st.dataframe(book[sheets[0]].head(8), use_container_width=True)

# ============================================================ 2 · FRAME
with tab_frame:
    st.header("Build or edit the code frame")
    if not st.session_state.raw:
        st.info("Load the data first.")
    else:
        q = st.selectbox("Question", list(st.session_state.raw), key="frame_q")
        st.caption(PATTERN_HELP)

        c1, c2, c3 = st.columns(3)
        with c1:
            if st.button("Start a blank frame", use_container_width=True):
                st.session_state.frames[q] = blank_frame()
        with c2:
            src = st.selectbox("Copy frame from", ["—"] + [k for k in st.session_state.frames if k != q],
                               label_visibility="collapsed")
            if src != "—" and st.button("Copy frame", use_container_width=True):
                st.session_state.frames[q] = [dict(r) for r in st.session_state.frames[src]]
        with c3:
            xl = st.file_uploader("Import frame from Excel", type=["xlsx"], key=f"fx_{q}",
                                  label_visibility="collapsed")
            if xl and st.button("Import", use_container_width=True):
                got = from_excel(BytesIO(xl.getvalue()))
                st.session_state.frames[q] = list(got.values())[0]

        frame = st.session_state.frames.get(q, blank_frame())
        grid = pd.DataFrame([{
            "NET": r["net"], "Code": r["code"], "Code label": r["label"],
            "Patterns": "\n".join(r["patterns"]),
        } for r in frame])

        edited = st.data_editor(
            grid, num_rows="dynamic", use_container_width=True, height=460,
            column_config={
                "Code": st.column_config.NumberColumn(width="small", format="%d"),
                "Code label": st.column_config.TextColumn(width="large"),
                "Patterns": st.column_config.TextColumn(width="large", help=PATTERN_HELP),
            }, key=f"editor_{q}")

        new_frame = []
        for _, r in edited.iterrows():
            if pd.isna(r["Code"]) or not str(r["Code"]).strip():
                continue
            new_frame.append({"net": r["NET"], "code": r["Code"],
                              "label": r["Code label"], "patterns": r["Patterns"]})
        if new_frame:
            new_frame = normalise(new_frame)
            st.session_state.frames[q] = new_frame

            issues = validate(new_frame)
            if issues:
                st.error("Fix before running:\n\n- " + "\n- ".join(issues))
            else:
                st.success(f"{len(new_frame)} codes, {len(set(r['net'] for r in new_frame))} NETs — frame is valid.")

        if st.session_state.frames:
            buf = BytesIO()
            to_excel(st.session_state.frames, buf)
            st.download_button("⬇️ Export all frames to Excel", buf.getvalue(),
                               file_name="code_frames.xlsx",
                               mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

# ============================================================ 3 · RUN
with tab_run:
    st.header("Run the coding")
    if not st.session_state.frames:
        st.info("Build a code frame first.")
    else:
        base_n = st.number_input("Base N for percentages (0 = use row count)", min_value=0,
                                 value=0, step=1)
        if st.button("▶️ Code all questions", type="primary"):
            for q in st.session_state.raw:
                if not st.session_state.frames.get(q):
                    st.warning(f"{q}: no frame, skipped.")
                    continue
                coded = coded_for(q)
                n = base_n or len(coded)
                freq = frequency_table(coded, st.session_state.frames[q], base_n=n)

                st.subheader(st.session_state.titles[q]["text"])
                unres = len(unresolved(coded, st.session_state.id_col))
                blank = sum(1 for cl in coded["Codes"] if 999 in cl)
                m1, m2, m3, m4 = st.columns(4)
                m1.metric("Responses", len(coded))
                m2.metric("Answered", len(coded) - blank)
                m3.metric("Codes used", int((freq["Count"] > 0).sum()))
                m4.metric("Needs review (996/997)", unres)

                show = freq[freq["Count"] > 0].sort_values("Count", ascending=False)
                st.dataframe(
                    show.style.format({"% of Total": "{:.1%}", "% of Answered": "{:.1%}"}),
                    use_container_width=True, height=320)

# ============================================================ 4 · REVIEW
with tab_review:
    st.header("Review and override")
    st.caption("Everything the rules could not place sits in 996 / 997. Fix those here — "
               "overrides are kept per respondent and survive a re-run of the rules.")
    if not st.session_state.frames:
        st.info("Run the coding first.")
    else:
        q = st.selectbox("Question", list(st.session_state.raw), key="rev_q")
        if st.session_state.frames.get(q):
            coded = coded_for(q)
            idc = st.session_state.id_col
            only_unres = st.checkbox("Show only 996 / 997", value=True)

            view = unresolved(coded, idc) if only_unres else pd.DataFrame({
                idc: coded[idc],
                "Verbatim": coded["Verbatim"],
                "Codes": coded["Codes"].map(lambda cl: "/".join(map(str, cl))),
            })
            view = view[view["Verbatim"].str.len() > 0]

            st.write(f"{len(view)} verbatim(s)")
            edited = st.data_editor(
                view, use_container_width=True, height=430, key=f"rev_{q}",
                column_config={
                    "Verbatim": st.column_config.TextColumn(width="large", disabled=True),
                    "Codes": st.column_config.TextColumn(
                        width="medium", help="Codes separated by / — e.g. 101/204/995"),
                })

            if st.button("Save overrides", type="primary", key=f"save_{q}"):
                store = st.session_state.overrides.setdefault(q, {})
                before = dict(zip(view[idc], view["Codes"]))
                saved = 0
                for _, r in edited.iterrows():
                    txt = str(r["Codes"]).strip()
                    if txt and txt != str(before.get(r[idc], "")):
                        try:
                            store[r[idc]] = [int(x) for x in txt.replace(",", "/").split("/") if x.strip()]
                            saved += 1
                        except ValueError:
                            st.error(f"{r[idc]}: '{txt}' is not a valid code list.")
                st.success(f"{saved} override(s) saved. {len(store)} in total for {q}.")

            if st.session_state.overrides.get(q):
                st.download_button(
                    "⬇️ Export overrides (.json)",
                    json.dumps({str(k): v for k, v in st.session_state.overrides[q].items()}, indent=2),
                    file_name=f"{q}_overrides.json", mime="application/json")

# ============================================================ 5 · EXPORT
with tab_export:
    st.header("Export the deliverable")
    if not st.session_state.frames:
        st.info("Nothing to export yet.")
    else:
        base_n = st.number_input("Base N (0 = row count)", min_value=0, value=0, step=1, key="exp_n")
        if st.button("📦 Build workbook", type="primary"):
            questions = {}
            for q in st.session_state.raw:
                if not st.session_state.frames.get(q):
                    continue
                questions[q] = {
                    "title": st.session_state.titles[q]["text"],
                    "frame": st.session_state.frames[q],
                    "coded": coded_for(q),
                    "id_col": st.session_state.id_col,
                    "extra_cols": st.session_state.extra_cols,
                }
            meta = {
                "title": st.session_state.meta.get("title", "Open-End Coding Deliverable"),
                "fields": {
                    "Job number": st.session_state.meta.get("job", ""),
                    "Client": st.session_state.meta.get("client", ""),
                    "Questions coded": ", ".join(questions),
                    "Total sample (N)": base_n or len(next(iter(questions.values()))["coded"]),
                },
            }
            buf = build_workbook(questions, meta, base_n=base_n or None)
            st.success("Workbook ready.")
            st.download_button(
                "⬇️ Download coded data + code structure",
                buf.getvalue(),
                file_name=f"{st.session_state.meta.get('job','project')}_OE_Coding_Final.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            st.caption("Formulas are written unevaluated by openpyxl. Excel recalculates them on open; "
                       "if you need cached values in the file itself, run scripts/recalc.py over it.")
