# -*- coding: utf-8 -*-
"""
oecoder.excel
Builds the client deliverable: Read Me, Code Structure, and per question a
coded-data sheet plus a 0/1 binary matrix for tabulation.
"""
from __future__ import annotations

from collections import Counter
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .frames import with_housekeeping

ARIAL = "Arial"
H_FILL = PatternFill("solid", fgColor="1F3864")
NET_FILL = PatternFill("solid", fgColor="D9E2F3")
Q_FILL = PatternFill("solid", fgColor="8EA9DB")
HK_FILL = PatternFill("solid", fgColor="F2F2F2")
H_FONT = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
B_FONT = Font(name=ARIAL, size=10, bold=True)
N_FONT = Font(name=ARIAL, size=10)
_THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = H_FONT
        cell.fill = H_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def _widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def build_workbook(questions: dict, meta: dict, base_n: int | None = None) -> BytesIO:
    """
    questions: {q_key: {"title": str, "frame": [...], "coded": DataFrame,
                        "id_col": str, "extra_cols": [str, ...]}}
               coded must carry 'Verbatim' and 'Codes' columns.
    meta:      free-form project details shown on the Read Me sheet.
    Returns an in-memory .xlsx ready to write to disk or hand to a download button.
    """
    wb = Workbook()

    # ----------------------------------------------------------- Read Me
    ws = wb.active
    ws.title = "Read Me"
    ws["A1"] = meta.get("title", "Open-End Coding Deliverable")
    ws["A1"].font = Font(name=ARIAL, size=14, bold=True, color="1F3864")

    lines = list(meta.get("fields", {}).items())
    lines += [
        ("", ""),
        ("Coding approach", "Frame built from a read of the verbatims and applied as keyword / regex rules, "
                            "with analyst overrides on reviewed cases. Multi-coding is allowed - one verbatim "
                            "can carry more than one code."),
        ("Base for %", "Percentages are on Total N. An 'Answered' base excluding blanks (code 999) is also shown."),
        ("Housekeeping codes", "995 Generic non-specific | 996 Any other mention | 997 Irrelevant / illegible | "
                               "998 Don't know / Can't say / NA | 999 No response"),
        ("", ""),
        ("Sheet guide", ""),
        ("  Code Structure", "The code frame for every question with counts and percentages - the sheet to "
                             "circulate for client sign-off."),
        ("  <Q> Coded", "Respondent-level data: ID, verbatim and the codes assigned."),
        ("  <Q> Binary", "The same data as a 0/1 matrix, one column per code, for the tabulation setup."),
    ]
    r = 3
    for k, v in lines:
        ws.cell(row=r, column=1, value=k).font = B_FONT
        c = ws.cell(row=r, column=2, value=v)
        c.font = N_FONT
        c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    _widths(ws, {"A": 26, "B": 105})

    # ----------------------------------------------------------- Code Structure
    ws = wb.create_sheet("Code Structure")
    ws["A1"] = "CODE STRUCTURE - " + meta.get("title", "")
    ws["A1"].font = Font(name=ARIAL, size=13, bold=True, color="1F3864")
    ws["A2"] = "Multi-code frame | Codes 995-999 are housekeeping codes"
    ws["A2"].font = Font(name=ARIAL, size=9, italic=True)

    hdr = ["NET", "Code", "Code label", "Count", "% of Total", "% of Answered"]
    r = 4
    for q, spec in questions.items():
        coded = spec["coded"]
        n = base_n or len(coded)
        answered = sum(1 for cl in coded["Codes"] if 999 not in cl)
        counts = Counter(c for cl in coded["Codes"] for c in cl)

        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        cell = ws.cell(row=r, column=1, value=spec.get("title", q))
        cell.font = Font(name=ARIAL, size=11, bold=True, color="1F3864")
        cell.fill = Q_FILL
        ws.row_dimensions[r].height = 20
        r += 1
        ws.cell(row=r, column=1,
                value=f"Base: Total N = {n}   |   Answered = {answered}   |   No response = {n - answered}"
                ).font = Font(name=ARIAL, size=9, italic=True)
        r += 1

        for i, h in enumerate(hdr, start=1):
            ws.cell(row=r, column=i, value=h)
        _header(ws, r, 6)
        r += 1

        last_net = None
        for row in with_housekeeping(spec["frame"]):
            code = int(row["code"])
            if row["net"] != last_net:
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
                c = ws.cell(row=r, column=1, value=str(row["net"]).upper())
                c.font = Font(name=ARIAL, size=10, bold=True, color="1F3864")
                c.fill = NET_FILL
                c.border = BORDER
                r += 1
                last_net = row["net"]
            ws.cell(row=r, column=2, value=code)
            ws.cell(row=r, column=3, value=row["label"])
            ws.cell(row=r, column=4, value=counts.get(code, 0))
            ws.cell(row=r, column=5, value=f"=IF($D{r}=\"\",\"\",$D{r}/{n})")
            ws.cell(row=r, column=6, value=f"=IF($D{r}=\"\",\"\",$D{r}/{max(answered,1)})")
            for c_ in range(1, 7):
                cell = ws.cell(row=r, column=c_)
                cell.font = N_FONT
                cell.border = BORDER
                if code >= 995:
                    cell.fill = HK_FILL
            ws.cell(row=r, column=2).alignment = Alignment(horizontal="center")
            ws.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(row=r, column=4).alignment = Alignment(horizontal="center")
            ws.cell(row=r, column=5).number_format = "0.0%"
            ws.cell(row=r, column=6).number_format = "0.0%"
            r += 1
        r += 2

    _widths(ws, {"A": 34, "B": 8, "C": 78, "D": 9, "E": 13, "F": 15})
    ws.freeze_panes = "A5"

    # ----------------------------------------------------------- per question
    for q, spec in questions.items():
        coded = spec["coded"]
        frame = with_housekeeping(spec["frame"])
        labels = {int(x["code"]): x["label"] for x in frame}
        id_col = spec["id_col"]
        extra = [c for c in spec.get("extra_cols", []) if c in coded.columns]
        mx = max((len(cl) for cl in coded["Codes"]), default=1)

        ws = wb.create_sheet(f"{q} Coded"[:31])
        ws["A1"] = spec.get("title", q)
        ws["A1"].font = Font(name=ARIAL, size=11, bold=True, color="1F3864")
        cols = [id_col] + extra + ["Verbatim response"] + [f"Code_{i}" for i in range(1, mx + 1)] \
               + ["No. of codes", "Code labels"]
        for i, h in enumerate(cols, start=1):
            ws.cell(row=3, column=i, value=h)
        _header(ws, 3, len(cols))

        vcol = 2 + len(extra)
        r = 4
        for _, row in coded.iterrows():
            cl = row["Codes"]
            ws.cell(row=r, column=1, value=row[id_col])
            for j, ec in enumerate(extra):
                ws.cell(row=r, column=2 + j, value=row[ec])
            ws.cell(row=r, column=vcol, value=row["Verbatim"])
            for i in range(mx):
                ws.cell(row=r, column=vcol + 1 + i, value=cl[i] if i < len(cl) else None)
            ws.cell(row=r, column=vcol + 1 + mx, value=len(cl))
            ws.cell(row=r, column=vcol + 2 + mx,
                    value=" | ".join(labels.get(c, "") for c in cl))
            for c_ in range(1, len(cols) + 1):
                cell = ws.cell(row=r, column=c_)
                cell.font = N_FONT
                cell.border = BORDER
                is_code = vcol + 1 <= c_ <= vcol + 1 + mx
                cell.alignment = Alignment(vertical="top",
                                           wrap_text=c_ in (vcol, len(cols)),
                                           horizontal="center" if is_code else "left")
            r += 1

        spec_w = {"A": 10, get_column_letter(vcol): 62}
        for j in range(len(extra)):
            spec_w[get_column_letter(2 + j)] = 16
        for i in range(mx):
            spec_w[get_column_letter(vcol + 1 + i)] = 8
        spec_w[get_column_letter(vcol + 1 + mx)] = 11
        spec_w[get_column_letter(vcol + 2 + mx)] = 60
        _widths(ws, spec_w)
        ws.freeze_panes = ws.cell(row=4, column=vcol + 1).coordinate
        ws.auto_filter.ref = f"A3:{get_column_letter(len(cols))}{r - 1}"

        # ------------------------------------------------------- binary
        counts = Counter(c for cl in coded["Codes"] for c in cl)
        used = [int(x["code"]) for x in frame if counts.get(int(x["code"]), 0) > 0]
        ws = wb.create_sheet(f"{q} Binary"[:31])
        ws.cell(row=1, column=1, value=f"{q} binary matrix (1 = code applies)").font = B_FONT
        for i, c in enumerate(used):
            cell = ws.cell(row=1, column=3 + i, value=labels.get(c, ""))
            cell.font = Font(name=ARIAL, size=7)
            cell.alignment = Alignment(text_rotation=90, vertical="bottom")
        ws.row_dimensions[1].height = 165

        bh = [id_col, "Verbatim"] + [f"C{c}" for c in used]
        for i, h in enumerate(bh, start=1):
            ws.cell(row=2, column=i, value=h)
        _header(ws, 2, len(bh))

        r = 3
        for _, row in coded.iterrows():
            s = set(row["Codes"])
            ws.cell(row=r, column=1, value=row[id_col]).font = N_FONT
            ws.cell(row=r, column=2, value=row["Verbatim"][:120]).font = N_FONT
            for i, c in enumerate(used):
                cell = ws.cell(row=r, column=3 + i, value=1 if c in s else 0)
                cell.font = N_FONT
                cell.alignment = Alignment(horizontal="center")
            r += 1
        ws.cell(row=r, column=2, value="TOTAL").font = B_FONT
        for i in range(len(used)):
            col = get_column_letter(3 + i)
            cell = ws.cell(row=r, column=3 + i, value=f"=SUM({col}3:{col}{r - 1})")
            cell.font = B_FONT
            cell.alignment = Alignment(horizontal="center")

        bw = {"A": 10, "B": 45}
        for i in range(len(used)):
            bw[get_column_letter(3 + i)] = 6
        _widths(ws, bw)
        ws.freeze_panes = "C3"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
